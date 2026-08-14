"""Builds the sales-by-employee .xlsx report, matching the manager's
existing template exactly (fonts, green header, alternating row fill,
number formats, formulas)."""

from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

FONT_NAME = "Times New Roman"
HEADER_FILL = PatternFill("solid", fgColor="FF00B050")
ALT_ROW_FILL = PatternFill("solid", fgColor="FFF2F2F2")
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFFFF")
NOTE_FONT = Font(name=FONT_NAME, size=9, color="FF595959")

FMT_MONEY = r"#\ ##0;\-#\ ##0;\—"
FMT_COUNT = r"0;;\—"
FMT_PERCENT = r"0%;\-0%;\—"
FMT_DELTA = r"\+#\ ##0;\-#\ ##0;\—"

HEADERS = [
    "Ф.И.О.",
    "Должность",
    "План",
    "Кол-во накладных",
    "Факт",
    "Средний чек",
    "Достижение %",
    "Отклонения/Перевып.",
]


def build_report(report_date: date, rows: list[dict], out_path: str) -> None:
    """rows: [{"name", "position", "plan", "count", "amount"}, ...]"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"

    ws["A1"] = f"Дата: {report_date.strftime('%d.%m.%Y')}"
    ws["A1"].font = Font(name=FONT_NAME, size=12, bold=True)

    for col, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    first_row = 3
    for i, row in enumerate(rows):
        r = first_row + i
        fill = ALT_ROW_FILL if i % 2 == 1 else None

        ws.cell(row=r, column=1, value=row["name"]).font = Font(name=FONT_NAME, size=11, bold=True)
        ws.cell(row=r, column=2, value=row["position"]).font = Font(name=FONT_NAME, size=11)

        c_plan = ws.cell(row=r, column=3, value=row["plan"])
        c_plan.font = Font(name=FONT_NAME, size=11, bold=True)
        c_plan.number_format = FMT_MONEY

        c_count = ws.cell(row=r, column=4, value=row["count"])
        c_count.font = Font(name=FONT_NAME, size=11)
        c_count.number_format = FMT_COUNT

        c_fact = ws.cell(row=r, column=5, value=row["amount"])
        c_fact.font = Font(name=FONT_NAME, size=11)
        c_fact.number_format = FMT_MONEY

        c_avg = ws.cell(row=r, column=6, value=f"=IF(D{r}=0,0,E{r}/D{r})")
        c_avg.font = Font(name=FONT_NAME, size=11)
        c_avg.number_format = FMT_MONEY

        c_pct = ws.cell(row=r, column=7, value=f"=IF(C{r}=0,0,E{r}/C{r})")
        c_pct.font = Font(name=FONT_NAME, size=11, bold=True)
        c_pct.number_format = FMT_PERCENT

        c_delta = ws.cell(row=r, column=8, value=f"=E{r}-C{r}")
        c_delta.font = Font(name=FONT_NAME, size=11)
        c_delta.number_format = FMT_DELTA

        if fill:
            for col in range(1, 9):
                ws.cell(row=r, column=col).fill = fill

    last_row = first_row + len(rows) - 1
    total_row = last_row + 1

    ws.cell(row=total_row, column=1, value="ИТОГО (без РОП)")
    not_rop = f"(B{first_row}:B{last_row}<>\"РОП\")"
    formulas = {
        3: f"=SUMPRODUCT({not_rop}*C{first_row}:C{last_row})",
        4: f"=SUMPRODUCT({not_rop}*D{first_row}:D{last_row})",
        5: f"=SUMPRODUCT({not_rop}*E{first_row}:E{last_row})",
        6: f"=IF(D{total_row}=0,0,E{total_row}/D{total_row})",
        7: f"=IF(C{total_row}=0,0,E{total_row}/C{total_row})",
        8: f"=E{total_row}-C{total_row}",
    }
    fmt_by_col = {3: FMT_MONEY, 4: FMT_COUNT, 5: FMT_MONEY, 6: FMT_MONEY, 7: FMT_PERCENT, 8: FMT_DELTA}
    for col in range(1, 9):
        cell = ws.cell(row=total_row, column=col, value=formulas.get(col))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        if col in fmt_by_col:
            cell.number_format = fmt_by_col[col]

    note_row = total_row + 2
    ws.cell(
        row=note_row,
        column=1,
        value=(
            f"Примечание: план — данные пользователя, кол-во накладных и факт — из UMAG на "
            f"{report_date.strftime('%d.%m.%Y')}. Средний чек = Факт/Кол-во накладных, "
            f"Достижение % = Факт/План, Отклонение = Факт-План — формулы."
        ),
    ).font = NOTE_FONT

    widths = [24, 12, 13, 16, 13, 13, 14, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(ord("A") + i - 1)].width = w

    wb.save(out_path)
